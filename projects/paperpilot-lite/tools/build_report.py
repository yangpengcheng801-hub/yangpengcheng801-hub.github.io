"""Build the PaperPilot course experiment report DOCX."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "PaperPilot_论文阅读助手_实验报告.docx"
UI_IMAGE = ROOT / "docs" / "assets" / "paperpilot-demo-viewport.png"
BENCHMARK = json.loads((ROOT / "docs" / "benchmark.json").read_text(encoding="utf-8"))

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BORDER = "CBD5E1"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=11, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(element, fill):
    shd = element.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        element.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths_dxa, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade(cell._tc.get_or_add_tcPr(), LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), 10, True, INK)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.alignment = (alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run_font(p.add_run(str(value)), 10)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(text), 9, False, MUTED)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    set_run_font(p.add_run(text), 11)
    return p


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_number(doc, text, num_id):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    set_run_font(p.add_run(text), 11)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph(style="Lead Callout")
    r = p.add_run(f"{label}  ")
    set_run_font(r, 10.5, True, DARK_BLUE)
    set_run_font(p.add_run(text), 10.5)
    return p


def add_code(doc, code):
    p = doc.add_paragraph(style="Code Block")
    for idx, line in enumerate(code.strip("\n").splitlines()):
        run = p.add_run(line)
        set_run_font(run, 9, name="Consolas")
        if idx < len(code.strip("\n").splitlines()) - 1:
            run.add_break()
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, 9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)

    callout = doc.styles.add_style("Lead Callout", 1)
    callout.font.name = "Calibri"
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    callout.font.size = Pt(10.5)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.line_spacing = 1.15
    shade(callout._element.get_or_add_pPr(), CALLOUT)

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(9)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.16)
    code.paragraph_format.line_spacing = 1.05
    shade(code._element.get_or_add_pPr(), LIGHT)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run("PaperPilot · 论文阅读助手实验报告"), 9, False, MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run("Page "), 9, False, MUTED)
    add_page_field(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(126)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_run_font(kicker.add_run("人工智能课程实验报告"), 11, True, BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_run_font(title.add_run("论文阅读助手 PaperPilot"), 30, True, INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    set_run_font(subtitle.add_run("基于本地文本分析与证据检索的课程智能体"), 15, False, DARK_BLUE)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(42)
    ppr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), BLUE)
    pbdr.append(bottom)
    ppr.append(pbdr)

    for line in ("姓名：__________________", "学号：__________________", "班级：__________________"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        set_run_font(p.add_run(line), 12)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(32)
    set_run_font(date.add_run("完成日期：2026 年 6 月 21 日"), 11, False, MUTED)
    doc.add_page_break()


def build():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "PaperPilot 论文阅读助手实验报告"
    doc.core_properties.subject = "人工智能课程智能体项目"
    doc.core_properties.author = "PaperPilot Project"
    add_cover(doc)

    doc.add_heading("摘  要", level=1)
    doc.add_paragraph(
        "本实验设计并实现了一个可在本地运行的论文阅读智能体 PaperPilot。系统面向学生和研究人员的首次论文阅读场景，"
        "提供 PDF、TXT 与 Markdown 解析、结构化摘要、科研要素抽取、概念关系图、批判性评审、证据式问答及阅读笔记导出。"
        "系统采用 Document、Analysis、Critic 与可选 Qwen Synthesis Agent 协作；本地算法构成稳定底座，配置阿里云百炼后增强综合表达，接口异常自动回退。实验通过语法编译、7 个单元测试、HTTP API 联调、真实浏览器交互和 12 题证据检索基准进行验证。"
        f"结果显示：单元测试 7/7 通过，12 个问题的目标证据全部进入返回段落，200 次分析的中位耗时为 {BENCHMARK['median_analysis_ms']} ms。"
    )
    add_callout(doc, "关键词", "论文阅读；智能体；文本摘要；关键词提取；证据检索；本地 Web 应用")

    doc.add_heading("项目交付物", level=2)
    add_table(
        doc,
        ["交付物", "文件/目录", "用途"],
        [
            ("程序代码", "app.py、paper_agent/、static/、tests/", "运行论文阅读助手与自动测试"),
            ("课程展示", "PaperPilot_论文阅读助手_课程展示.pptx", "课堂讲解项目背景、实现与实验"),
            ("实验报告", "PaperPilot_论文阅读助手_实验报告.docx", "记录需求、设计、实现、测试与结论"),
        ],
        [1700, 3900, 3760],
    )
    doc.add_heading("1 实验目的", level=1)
    add_bullet(doc, "理解智能体任务分解方法，将论文阅读拆分为解析、理解、检索、回答和导出等步骤。")
    add_bullet(doc, "掌握 Python 文本处理、PDF 读取、本地 HTTP 服务与浏览器前端的组合实现。")
    add_bullet(doc, "建立“回答必须可回到原文核验”的证据意识，降低脱离论文内容作答的风险。")
    add_bullet(doc, "通过自动测试与可复现实验评价系统的功能正确性、检索效果和运行效率。")

    doc.add_heading("2 需求分析", level=1)
    doc.add_heading("2.1 用户场景", level=2)
    doc.add_paragraph(
        "目标用户需要在短时间内判断一篇论文是否值得精读，并快速定位研究问题、方法、贡献和实验结果。"
        "如果系统只给出一段流畅摘要，用户仍难以判断结论来自何处。因此，本项目把“证据位置”作为问答输出的必要组成。"
    )
    doc.add_heading("2.2 功能需求", level=2)
    add_table(
        doc,
        ["编号", "需求", "验收标准"],
        [
            ("F1", "导入 PDF/TXT/Markdown 或粘贴正文", "能提取不少于 80 个字符的有效文本"),
            ("F2", "生成结构化阅读结果", "包含摘要、科研要素、概念图、评审和章节"),
            ("F3", "基于论文内容回答问题", "返回答案以及至少 1 个证据段落编号"),
            ("F4", "导出阅读笔记", "浏览器下载 Markdown 文件"),
            ("F5", "双模式可运行", "本地离线可用；百炼增强失败自动回退"),
        ],
        [780, 3600, 4980],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_heading("2.3 非功能需求", level=2)
    add_bullet(doc, "可复现：相同输入在相同版本下得到稳定结果。")
    add_bullet(doc, "可检查：问答展示原文证据编号；重要结论提示用户回到原文核对。")
    add_bullet(doc, "易部署：除 PDF 解析库 pypdf 外，Web 服务仅使用 Python 标准库。")
    add_bullet(doc, "易演示：内置示例论文，断网环境下仍可完成主要流程。")

    doc.add_page_break()
    doc.add_heading("3 系统总体设计", level=1)
    doc.add_heading("3.1 分层架构", level=2)
    architecture_num = create_decimal_numbering(doc)
    add_number(doc, "输入层接收文件字节或用户粘贴的文本，并完成格式判断与编码处理。", architecture_num)
    add_number(doc, "Document Agent 解析段落与章节，Analysis Agent 抽取摘要、关键词与科研要素。", architecture_num)
    add_number(doc, "Critic Agent 构建概念关系、归纳优势/局限/追问；可选 Qwen Agent 在证据约束下增强表达。", architecture_num)
    add_number(doc, "交互输出层展示结构化结果、Agent 轨迹、证据问答并支持笔记导出。", architecture_num)
    add_callout(doc, "核心原则", "速度不能以牺牲可追溯性为代价；摘要尽量来自原文，回答附带段落级证据。")

    doc.add_heading("3.2 数据流", level=2)
    add_table(
        doc,
        ["阶段", "输入", "处理", "输出"],
        [
            ("解析", "PDF/TXT/MD", "读取、解码、规范化空白", "纯文本"),
            ("理解", "纯文本", "段落/句子/章节识别", "PaperDocument"),
            ("分析", "PaperDocument", "关键词与句子评分、信号句抽取", "结构化结果"),
            ("问答", "问题 + 段落", "词项重合度和章节信号排序", "答案 + 证据"),
            ("导出", "分析结果", "Markdown 模板组织", "阅读笔记"),
        ],
        [1150, 2050, 3560, 2600],
    )

    doc.add_heading("3.3 项目结构", level=2)
    add_code(doc, """paper-reading-assistant/
├─ app.py                 # HTTP 服务与 API
├─ paper_agent/core.py    # 本地解析、图谱、评审和问答
├─ paper_agent/llm.py     # 阿里云百炼兼容客户端
├─ paper_agent/orchestrator.py # 多智能体编排与回退
├─ static/                # HTML / CSS / JavaScript 界面
├─ samples/               # 示例论文
├─ tests/                 # 单元测试
├─ tools/                 # 截图、基准与报告生成工具
└─ deliverables/          # PPT 与实验报告""")

    doc.add_heading("4 关键实现", level=1)
    doc.add_heading("4.1 文档解析与结构识别", level=2)
    doc.add_paragraph(
        "PDF 文件使用 pypdf 按页提取文本；TXT/Markdown 依次尝试 UTF-8-SIG、UTF-8 和 GB18030 编码。"
        "规范化后，系统通过空行和标点切分段落/句子，并匹配 Abstract、Introduction、Method、Experiments、Results、Conclusion 等常见标题。"
    )
    add_code(doc, """document = PaperDocument(filename=filename, text=text)
analysis = agent.analyze(document)
documents[document.doc_id] = document""")

    doc.add_heading("4.2 摘要与关键词", level=2)
    doc.add_paragraph(
        "关键词模块分别处理英文单词和中文连续字符，过滤停用词后按频次、词长和特异性排序。"
        "摘要模块对句子计算关键词重合度、位置权重、长度得分和科研信号词加成，再按原文顺序返回高分句，避免摘要顺序混乱。"
    )
    add_code(doc, """score = (keyword_overlap + 0.6 * length_score) \
        * position_bonus * scientific_signal
top_sentences = sorted(scored_sentences, key=score, reverse=True)""")

    doc.add_heading("4.3 证据式问答", level=2)
    doc.add_paragraph(
        "问题和段落经过相同的分词函数。检索分数由完全词项重合、模糊包含和方法/结果/结论章节信号构成。"
        "系统返回分数最高的 3 个段落，并在答案后标注 Pn 证据编号。该方法不等同于深层语义推理，但具有速度快、结果稳定和易核验的优点。"
    )
    add_code(doc, """score = overlap + fuzzy_overlap * 0.8 + section_signal
evidence.sort(key=lambda item: item.score, reverse=True)
answer = sentence + " [证据P%d]" % paragraph_id""")

    doc.add_heading("4.4 多智能体编排与百炼增强", level=2)
    doc.add_paragraph(
        "编排器首先运行本地分析，随后在检测到 DASHSCOPE_API_KEY 时调用百炼 OpenAI 兼容接口。"
        "模型提示明确把论文视为不可信数据，并要求严格 JSON、不得虚构数字。问答阶段先由本地检索选出证据，再让 Qwen 仅根据这些段落作答；任何网络或接口错误都会保留本地结果。"
    )
    add_code(doc, """analysis = local_agent.analyze(document)
if bailian.available:
    analysis.update(bailian.enhance_analysis(document.text, analysis))
# 失败时保留 analysis，并记录 fallback 轨迹""")

    doc.add_page_break()
    doc.add_heading("5 程序界面与操作流程", level=1)
    if UI_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        picture = p.add_run().add_picture(str(UI_IMAGE), width=Inches(6.35))
        picture._inline.docPr.set("descr", "PaperPilot 本地网页界面，展示论文导入、结构化摘要和关键词")
        picture._inline.docPr.set("title", "PaperPilot 程序界面")
        add_caption(doc, "图 1  PaperPilot 本地 Web 界面：论文导入、结构化结果与关键词")
    doc.add_heading("5.1 启动步骤", level=2)
    startup_num = create_decimal_numbering(doc)
    add_number(doc, r"进入项目目录：D:\APP\pycharm\python learn\paper-reading-assistant。", startup_num)
    add_number(doc, "执行 python -m pip install -r requirements.txt 安装 PDF 解析依赖。", startup_num)
    add_number(doc, "执行 python app.py，并在浏览器访问 http://127.0.0.1:8765。", startup_num)
    add_number(doc, "上传论文或点击“加载示例”，再点击“开始阅读”。", startup_num)
    add_number(doc, "查看摘要、概念图、批判性评审与 Agent 轨迹；输入问题得到证据式回答。", startup_num)
    add_number(doc, "如需大模型增强，运行 start_llm.bat 并在提示时输入百炼密钥。", startup_num)

    doc.add_page_break()
    doc.add_heading("5.2 API 设计", level=2)
    add_table(
        doc,
        ["接口", "方法", "作用", "关键返回值"],
        [
            ("/api/health", "GET", "服务健康检查", "mode、llm_enabled、model"),
            ("/api/sample", "GET", "读取示例论文", "filename、text"),
            ("/api/analyze", "POST", "分析上传文件/文本", "analysis、doc_id"),
            ("/api/ask", "POST", "检索论文并回答", "answer、evidence"),
            ("/api/export", "POST", "生成 Markdown 笔记", "filename、markdown"),
        ],
        [1900, 900, 3060, 3500],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("6 实验方法", level=1)
    doc.add_heading("6.1 实验环境", level=2)
    add_table(
        doc,
        ["项目", "配置"],
        [
            ("操作系统", "Windows，PowerShell"),
            ("语言", "Python 3，JavaScript"),
            ("主要依赖", "pypdf；python-docx（报告生成）"),
            ("浏览器验收", "Google Chrome 无头模式"),
            ("服务地址", "127.0.0.1:8765"),
        ],
        [2200, 7160],
    )

    doc.add_heading("6.2 功能正确性测试", level=2)
    doc.add_paragraph("使用 unittest 对智能体核心进行 7 项自动化测试。测试命令为：")
    add_code(doc, "python -m unittest discover -s tests -v")
    add_table(
        doc,
        ["测试项", "断言", "结果"],
        [
            ("标题与章节识别", "识别 PaperPilot 标题及摘要/方法/结果章节", "通过"),
            ("分析字段完整性", "摘要、关键词、贡献、方法、结果均非空", "通过"),
            ("图谱与评审", "概念节点、优势与追问均非空，Agent 轨迹完整", "通过"),
            ("离线回退", "未提供密钥时编排器稳定返回本地结果", "通过"),
            ("问答证据", "命中包含 86.0% 的实验结果段落", "通过"),
            ("笔记导出", "Markdown 包含研究贡献与实验结果章节", "通过"),
            ("PDF 文本提取", "从示例 PDF 提取标题且文本超过 3,000 字符", "通过"),
        ],
        [2700, 5100, 1560],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    doc.add_heading("6.3 证据检索基准", level=2)
    doc.add_paragraph(
        "基准脚本 tools/benchmark.py 针对内置示例论文设计 12 个可核验问题，并为每个问题指定期望事实。"
        "若返回的前 3 个证据段落包含该事实，则记为命中。性能测试重复执行完整分析 200 次，统计中位数和 P95。"
    )
    add_callout(doc, "边界说明", "该基准用于验证项目的可复现功能，不代表在任意学科、任意论文上的通用准确率。")

    doc.add_heading("7 实验结果", level=1)
    add_table(
        doc,
        ["指标", "结果", "判定"],
        [
            ("Python 语法编译", "通过", "无语法错误"),
            ("单元测试", "7/7 通过", "核心功能正确"),
            ("API 端到端联调", "健康检查、分析、问答均通过", "服务链路可用"),
            ("浏览器交互", "加载、分析、提问、显示证据均通过", "界面可演示"),
            ("证据问题命中", f"{BENCHMARK['evidence_hits']}/{BENCHMARK['question_cases']}（{BENCHMARK['evidence_hit_rate_percent']}%）", "样例基准全部命中"),
            ("分析耗时中位数", f"{BENCHMARK['median_analysis_ms']} ms", "满足本地交互需求"),
            ("分析耗时 P95", f"{BENCHMARK['p95_analysis_ms']} ms", "波动较小"),
        ],
        [2600, 3600, 3160],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_heading("7.1 结果分析", level=2)
    add_bullet(doc, "本地抽取式分析在 4,824 字符的示例论文上可在毫秒级完成，课堂演示响应足够快。")
    add_bullet(doc, "12 个事实型问题均返回包含目标事实的证据段落，说明词项重合检索对显式事实定位有效。")
    add_bullet(doc, "自动测试、HTTP API 与浏览器三层结果一致，证明核心算法、服务层与前端交互已连通。")
    add_bullet(doc, "系统默认无需付费 API；百炼为可选增强且具备自动回退，兼顾效果与展示稳定性。")

    doc.add_heading("8 讨论与局限", level=1)
    doc.add_heading("8.1 优点", level=2)
    add_bullet(doc, "可解释：关键词、摘要和问答都可追溯到原文内容。")
    add_bullet(doc, "可复现：核心算法确定性强，基准与测试脚本均随项目交付。")
    add_bullet(doc, "依赖轻：无需数据库和前端框架，便于在 PyCharm 或命令行直接运行。")
    add_bullet(doc, "功能可解释：概念图、Critic Agent 和协作轨迹使分析过程不再是黑盒。")

    doc.add_heading("8.2 局限", level=2)
    add_bullet(doc, "扫描版 PDF 不含可提取文本，需要先进行 OCR。")
    add_bullet(doc, "复杂双栏排版可能导致文本顺序与视觉阅读顺序不一致。")
    add_bullet(doc, "基于词项重合的检索对同义改写、跨段推理、公式和图表理解能力有限。")
    add_bullet(doc, "抽取式摘要可核验但语言组织能力弱于大型语言模型。")

    doc.add_heading("8.3 改进方向", level=2)
    improvement_num = create_decimal_numbering(doc)
    add_number(doc, "加入 OCR 模块与版面分析，改善扫描件和多栏论文的解析质量。", improvement_num)
    add_number(doc, "增加向量嵌入检索，提高同义表达和中英文交叉提问的召回能力。", improvement_num)
    add_number(doc, "扩展多论文横向对比、引用网络和长期知识库，在当前 Qwen 证据增强基础上支持系统综述。", improvement_num)
    add_number(doc, "扩展图表、公式、参考文献与引用网络分析，实现多模态论文阅读。", improvement_num)

    doc.add_heading("9 实验结论", level=1)
    doc.add_paragraph(
        "本实验完成了 PaperPilot 论文阅读助手的需求分析、系统设计、代码实现、界面开发与验证。"
        "系统能够在本地将论文转换为结构化阅读结果，并通过段落检索支持证据式问答。"
        f"自动测试 7/7 通过，样例基准 {BENCHMARK['question_cases']} 个问题全部命中目标证据，分析中位耗时 {BENCHMARK['median_analysis_ms']} ms。"
        "结果说明：本地确定性底座与可选大模型增强可以同时满足课堂稳定性、分析质量和证据可追溯性。"
    )
    add_callout(doc, "最终结论", "项目已形成“可运行程序 + 可展示 PPT + 可复现实验报告”的完整课程成果。")

    doc.add_heading("附录 A 复现命令", level=1)
    add_code(doc, r"""cd /d "D:\APP\pycharm\python learn\paper-reading-assistant"
python -m pip install -r requirements.txt
python -m unittest discover -s tests -v
python tools\benchmark.py
python app.py
rem 或运行 start_llm.bat 启用百炼增强""")
    doc.add_heading("附录 B 主要文件说明", level=2)
    add_table(
        doc,
        ["文件", "说明"],
        [
            ("paper_agent/core.py", "智能体核心算法与数据结构"),
            ("paper_agent/llm.py", "百炼 OpenAI 兼容客户端，密钥仅从环境读取"),
            ("paper_agent/orchestrator.py", "多智能体编排、增强与自动回退"),
            ("app.py", "本地 HTTP 服务、API 路由和会话管理"),
            ("static/index.html", "单页应用结构"),
            ("static/styles.css", "界面视觉与响应式布局"),
            ("static/app.js", "文件读取、API 调用、结果渲染与导出"),
            ("tests/test_agent.py", "核心单元测试"),
            ("tools/benchmark.py", "证据命中与性能基准"),
        ],
        [3200, 6160],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
